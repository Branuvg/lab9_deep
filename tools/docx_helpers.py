"""Utilidades para insertar respuestas con ecuaciones nativas (OMML) en el docx del laboratorio."""
import re, copy
from lxml import etree
import latex2mathml.converter
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

XSL = etree.XSLT(etree.parse(r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL"))
BLUE = RGBColor(0x1F, 0x3A, 0x93)
M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"


def omml(latex):
    """LaTeX -> elemento <m:oMath>."""
    # \quad y \qquad se pierden en la conversión: se sustituyen por espacios em (U+2003) dentro de \text{}
    latex = latex.replace(r"\qquad", "\\text{\u2003\u2003}").replace(r"\quad", "\\text{\u2003}")
    mml = latex2mathml.converter.convert(latex)
    tree = XSL(etree.fromstring(mml))
    return copy.deepcopy(tree.getroot())


def style_run(r, bold=False, italic=False, size=11, color=None):
    rpr = r._r.get_or_add_rPr()
    f = OxmlElement("w:rFonts")
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme"):
        f.set(qn(a), "minorBidi")
    rpr.insert(0, f)
    r.bold = bold; r.italic = italic; r.font.size = Pt(size)
    if color is not None: r.font.color.rgb = color


def rich(p, text, size=11, italic=False):
    """Agrega a p texto con **negrita** y $latex$ en línea."""
    bold = False
    for seg in re.split(r"(\$[^$]+\$)", text):
        if not seg: continue
        if seg.startswith("$") and seg.endswith("$"):
            p._p.append(omml(seg[1:-1])); continue
        for k, part in enumerate(seg.split("**")):
            if k > 0: bold = not bold
            if part: style_run(p.add_run(part), bold=bold, italic=italic, size=size)


class Cursor:
    """Inserta bloques consecutivos a partir de un elemento ancla del documento."""
    def __init__(self, doc, anchor, figdir):
        self.doc, self.el, self.figdir = doc, getattr(anchor, "_p", anchor), figdir

    def _move(self, new_el):
        self.el.addnext(new_el); self.el = new_el

    def para(self, text="", label=None, italic=False, size=11, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=False):
        p = self.doc.add_paragraph()
        pf = p.paragraph_format
        if indent: pf.left_indent = Inches(0.25)
        pf.space_before = Pt(6); pf.space_after = Pt(6); pf.line_spacing = 1.15
        if label: style_run(p.add_run(label + " "), bold=True, size=size, color=BLUE)
        rich(p, text, size=size, italic=italic)
        if align is not None: p.alignment = align
        self._move(p._p); return p

    def eq(self, latex):
        p = self.doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(6)
        para = etree.SubElement(p._p, f"{{{M_NS}}}oMathPara")
        para.append(omml(latex))
        self._move(p._p); return p

    def image(self, name, width=6.0, caption=None):
        p = self.doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(f"{self.figdir}/{name}", width=Inches(width)); self._move(p._p)
        if caption: self.para(caption, italic=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)

    def table(self, header, rows, caption=None):
        t = self.doc.add_table(rows=1 + len(rows), cols=len(header))
        for j, h in enumerate(header):
            c = t.rows[0].cells[j]; c.text = ""; rich(c.paragraphs[0], f"**{h}**", size=9.5)
            shade(c, "D9E2F3")
        for i, row in enumerate(rows):
            for j, v in enumerate(row):
                c = t.rows[i + 1].cells[j]; c.text = ""; rich(c.paragraphs[0], str(v), size=9.5)
        borders(t); t.alignment = 1
        self._move(t._tbl)
        if caption: self.para(caption, italic=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
        else: self.para("", indent=False)


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    s = OxmlElement("w:shd"); s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), hexcolor)
    tcPr.append(s)


def borders(table):
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:color"), "808080")
        b.append(e)
    tblPr.append(b)


def find(doc, prefix):
    for p in doc.paragraphs:
        if p.text.strip().startswith(prefix):
            return p
    raise KeyError(prefix)
